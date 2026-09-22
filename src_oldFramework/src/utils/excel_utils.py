import os
import openpyxl
from utils.config import TESTDATA_DIR

class ExcelUtil:
    @staticmethod
    def get_sheet(file_name: str, sheet_name: str) -> list[dict[str, str]]:
        """
        Reads an Excel sheet and returns data as a list of dictionaries.
        Each row is represented as a dict with column headers as keys.

        :param file_name: Excel file name (without extension, expects .xlsx)
        :param sheet_name: Name of the sheet to read
        :return: List of dicts [{header:value}, ...]
        """
        data = []
        file_path = os.path.join(TESTDATA_DIR, f"{file_name}.xlsx")

        try:
            workbook = openpyxl.load_workbook(file_path)
            sheet = workbook[sheet_name]

            # Read header row (first row)
            headers = [
                str(sheet.cell(row=1, column=col).value).strip()
                for col in range(1, sheet.max_column + 1)
            ]

            # Loop through rows starting from row 2
            for row in range(2, sheet.max_row + 1):
                row_dict = {}
                for col, header in enumerate(headers, start=1):
                    cell_value = sheet.cell(row=row, column=col).value
                    row_dict[header] = "" if cell_value is None else str(cell_value)
                data.append(row_dict)

            workbook.close()

        except Exception as e:
            print(f"[ExcelUtil] Error reading Excel file: {e}")

        return data
